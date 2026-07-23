import json
import os
import threading
import time
import random
from datetime import datetime

import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog, scrolledtext

from education_system.post_18.university_system.infrastructure.database.db import DEFAULT_DB_PATH, sqlite3
from education_system.post_18.university_system.infrastructure.auth import UserAuth
from education_system.post_18.university_system.infrastructure.shared_context import get_auth

try:
    from education_system.post_18.university_system.infrastructure.ai.ai_detector.detector import AIDetector
    _AI_DETECTOR_IMPORT_ERROR = None
except Exception as import_error:
    AIDetector = None
    _AI_DETECTOR_IMPORT_ERROR = import_error

try:
    import textract
    TEXTRACT_AVAILABLE = True
except ImportError:
    TEXTRACT_AVAILABLE = False

try:
    from pypdf import PdfReader
    PYPDF2_AVAILABLE = True
except ImportError:
    PYPDF2_AVAILABLE = False

try:
    import docx
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False

from education_system.post_18.university_system.core.i18n import get_text, _

def create_analysis_view(self, parent):
    """Create the main text analysis view"""
    analysis_frame = ttk.Frame(parent)
    analysis_frame.pack(fill='both', expand=True)

    # Create two-column layout
    left_frame = ttk.Frame(analysis_frame)
    left_frame.pack(side='left', fill='both', expand=True, padx=(0, 5))

    right_frame = ttk.Frame(analysis_frame)
    right_frame.pack(side='right', fill='both', expand=True, padx=(5, 0))

    # Left side - Input
    self.create_input_section(left_frame)

    # Right side - Results
    self.create_results_section(right_frame)


def create_input_section(self, parent):
    """Create text input section"""
    # Input card
    input_card = ttk.Frame(parent, style='Card.TFrame')
    input_card.pack(fill='both', expand=True, padx=5, pady=5)

    # Card title
    ttk.Label(input_card, text="Text Input", style='Title.TLabel').pack(anchor='w', padx=15, pady=(15, 5))

    # Metadata inputs
    metadata_frame = ttk.Frame(input_card)
    metadata_frame.pack(fill='x', padx=15, pady=(0, 10))

    # Row 1: Student and Title
    row1 = ttk.Frame(metadata_frame)
    row1.pack(fill='x', pady=(0, 5))

    ttk.Label(row1, text="Student:").pack(side='left')
    self.student_var = tk.StringVar()
    self.student_combo = ttk.Combobox(row1, textvariable=self.student_var, width=25, state='readonly')
    self.student_combo.pack(side='left', padx=(5, 15))

    ttk.Label(row1, text="Title:").pack(side='left')
    self.title_var = tk.StringVar()
    self.title_entry = ttk.Entry(row1, textvariable=self.title_var, width=30)
    self.title_entry.pack(side='left', padx=(5, 0))

    # Row 2: Module and Assignment
    row2 = ttk.Frame(metadata_frame)
    row2.pack(fill='x')

    ttk.Label(row2, text="Module:").pack(side='left')
    self.module_var = tk.StringVar()
    self.module_combo = ttk.Combobox(row2, textvariable=self.module_var, width=20, state='readonly')
    self.module_combo.pack(side='left', padx=(5, 15))
    self.module_combo.bind('<<ComboboxSelected>>', self._on_module_selected)

    ttk.Label(row2, text="Assignment:").pack(side='left')
    self.assignment_var = tk.StringVar()
    self.assignment_combo = ttk.Combobox(row2, textvariable=self.assignment_var, width=25, state='readonly')
    self.assignment_combo.pack(side='left', padx=(5, 0))

    # Load data from database
    self._load_database_data()

    # Text input area
    ttk.Label(input_card, text="Text to Analyze:", style='Subtitle.TLabel').pack(anchor='w', padx=15, pady=(10, 5))

    # Buttons and word count packed BEFORE text area so they stay visible
    button_frame = ttk.Frame(input_card)
    button_frame.pack(side='bottom', fill='x', padx=15, pady=(0, 15))

    self.upload_btn = ttk.Button(button_frame, text="📁 Load File", command=self.load_file)
    self.upload_btn.pack(side='left', padx=(0, 10))

    self.clear_btn = ttk.Button(button_frame, text="🗑️ Clear", command=self.clear_input)
    self.clear_btn.pack(side='left', padx=(0, 10))

    self.analyze_btn = ttk.Button(
        button_frame,
        text="🔍 Analyze Text",
        command=self.analyze_text,
        style='Accent.TButton'
    )
    self.analyze_btn.pack(side='right')

    self.word_count_label = ttk.Label(input_card, text="Words: 0", style='Subtitle.TLabel')
    self.word_count_label.pack(side='bottom', anchor='e', padx=15, pady=(0, 5))

    # Text widget fills remaining space
    text_frame = ttk.Frame(input_card)
    text_frame.pack(fill='both', expand=True, padx=15, pady=(0, 10))

    self.text_input = scrolledtext.ScrolledText(
        text_frame,
        wrap=tk.WORD,
        width=50,
        height=20,
        bg=self.colors['bg_secondary'],
        fg=self.colors['text_primary'],
        insertbackground=self.colors['text_primary'],
        selectbackground=self.colors['accent']
    )
    self.text_input.pack(fill='both', expand=True)

    self.text_input.bind('<KeyRelease>', self.update_word_count)


def create_results_section(self, parent):
    """Create results display section"""
    # Results card
    results_card = ttk.Frame(parent, style='Card.TFrame')
    results_card.pack(fill='both', expand=True, padx=5, pady=5)

    # Card title
    ttk.Label(results_card, text="Analysis Results", style='Title.TLabel').pack(anchor='w', padx=15, pady=(15, 5))

    # Results container
    self.results_container = ttk.Frame(results_card)
    self.results_container.pack(fill='both', expand=True, padx=15, pady=15)

    # Initial empty state
    self.show_empty_results()


def show_empty_results(self):
    """Show empty state for results"""
    for widget in self.results_container.winfo_children():
        widget.destroy()

    empty_label = ttk.Label(self.results_container,
                          text="Enter text and click 'Analyze Text' to see results",
                          style='Subtitle.TLabel')
    empty_label.pack(expand=True)


def load_file(self):
    """Load text from file (supports TXT, PDF, DOCX)"""
    # Build file types list based on available libraries
    filetypes = [("Text files", "*.txt")]

    if PYPDF2_AVAILABLE or TEXTRACT_AVAILABLE:
        filetypes.append(("PDF files", "*.pdf"))

    if PYTHON_DOCX_AVAILABLE or TEXTRACT_AVAILABLE:
        filetypes.append(("Word documents", "*.docx"))

    filetypes.append(("All files", "*.*"))

    file_path = filedialog.askopenfilename(
        title="Select document file",
        filetypes=filetypes
    )

    if file_path:
        try:
            content = self._extract_text_from_file(file_path)

            if content:
                self.text_input.delete('1.0', tk.END)
                self.text_input.insert('1.0', content)
                self.update_word_count()

                # Auto-fill title from filename if empty
                filename = os.path.basename(file_path)
                if not self.title_var.get():
                    filename_without_ext = os.path.splitext(filename)[0]
                    self.title_var.set(filename_without_ext)

                self.update_status(f"Loaded file: {filename}")
            else:
                messagebox.showwarning("Warning", "No text could be extracted from the file.")

        except Exception as e:
            messagebox.showerror("Error", f"Failed to load file: {str(e)}")


def clear_input(self):
    """Clear all input fields"""
    self.text_input.delete('1.0', tk.END)
    self.student_var.set("")
    self.title_var.set("")
    self.module_var.set("")
    self.assignment_var.set("")
    self.update_word_count()
    self.show_empty_results()


def analyze_text(self):
    """Analyze the input text"""
    text = self.text_input.get('1.0', tk.END).strip()

    if not text:
        messagebox.showwarning("Warning", "Please enter text to analyze")
        return

    if len(text) < 50:
        messagebox.showwarning("Warning", "Text is too short for meaningful analysis (minimum 50 characters)")
        return

    # Update UI state
    self.analyze_btn.config(state='disabled', text='🔄 Analyzing...')
    self.update_status("Analyzing text...")
    self.show_progress(True)

    # Run analysis in background thread
    def analyze_thread():
        try:
            # Get metadata
            student_id = self.student_var.get() or None
            title = self.title_var.get() or None
            course_code = self.module_var.get() or None
            assignment_id = self.assignment_var.get() or None

            # Run analysis
            result = self.detector.analyze_text_enhanced(
                text=text,
                title=title,
                student_id=student_id,
                course_code=course_code,
                assignment_id=assignment_id
            )

            # Store results
            self.analysis_results = result
            self.current_submission_id = result.get('submission_id')

            # Update UI in main thread
            self.root.after(0, self.display_results, result)

        except Exception as e:
            self.root.after(0, self.analysis_error, str(e))

    threading.Thread(target=analyze_thread, daemon=True).start()


def display_results(self, result):
    """Display analysis results"""
    # Clear previous results
    for widget in self.results_container.winfo_children():
        widget.destroy()

    # Main results frame
    main_results = ttk.Frame(self.results_container)
    main_results.pack(fill='both', expand=True)

    # AI Score display
    self.create_score_display(main_results, result)

    # Detailed analysis
    self.create_detailed_analysis(main_results, result)

    # Recommendations
    self.create_recommendations(main_results, result)

    # Reset UI state
    self.analyze_btn.config(state='normal', text='🔍 Analyze Text')
    self.show_progress(False)
    self.update_status("Analysis complete")

    # Refresh statistics and history
    self.refresh_statistics()
    self.refresh_history()


def create_score_display(self, parent, result):
    """Create main score display"""
    score_frame = ttk.Frame(parent, style='Card.TFrame')
    score_frame.pack(fill='x', pady=(0, 10))

    # AI Score
    ai_score = result.get('ai_score', 0)
    confidence = result.get('confidence', 0)
    is_ai_generated = result.get('is_ai_generated', False)

    # Score header
    header_frame = ttk.Frame(score_frame)
    header_frame.pack(fill='x', padx=15, pady=15)

    # Risk indicator
    risk_color = self.get_risk_color(ai_score)
    risk_text = self.get_risk_text(ai_score)

    risk_indicator = tk.Label(header_frame, text="●", font=('Arial', 20),
                            fg=risk_color, bg=self.colors['bg_tertiary'])
    risk_indicator.pack(side='left')

    ttk.Label(header_frame, text=f"{risk_text} Risk",
             font=('Segoe UI', 14, 'bold')).pack(side='left', padx=(10, 0))

    # Score display
    score_display_frame = ttk.Frame(score_frame)
    score_display_frame.pack(fill='x', padx=15, pady=(0, 15))

    # AI Score
    score_label = ttk.Label(score_display_frame, text=f"{ai_score:.1%}",
                           font=('Segoe UI', 32, 'bold'))
    score_label.pack(side='left')

    # Details
    details_frame = ttk.Frame(score_display_frame)
    details_frame.pack(side='left', padx=(20, 0), fill='y')

    ttk.Label(details_frame, text="AI Generated Probability",
             style='Subtitle.TLabel').pack(anchor='w')
    ttk.Label(details_frame, text=f"Confidence: {confidence:.1%}",
             style='Subtitle.TLabel').pack(anchor='w')
    ttk.Label(details_frame, text=f"Status: {'AI Generated' if is_ai_generated else 'Human Written'}",
             style='Subtitle.TLabel').pack(anchor='w')


def create_detailed_analysis(self, parent, result):
    """Create detailed analysis section"""
    details_frame = ttk.Frame(parent, style='Card.TFrame')
    details_frame.pack(fill='both', expand=True, pady=(0, 10))

    ttk.Label(details_frame, text="Detailed Analysis",
             style='Title.TLabel').pack(anchor='w', padx=15, pady=(15, 10))

    # Analysis notebook
    analysis_notebook = ttk.Notebook(details_frame)
    analysis_notebook.pack(fill='both', expand=True, padx=15, pady=(0, 15))

    # Pattern Analysis tab
    self.create_pattern_tab(analysis_notebook, result)

    # Sentence Analysis tab (if available)
    if result.get('sentence_analysis'):
        self.create_sentence_tab(analysis_notebook, result)

    # Advanced Analysis tab (if available)
    if result.get('advanced_analyses'):
        self.create_advanced_analysis_tab(analysis_notebook, result)


def create_pattern_tab(self, notebook, result):
    """Create pattern analysis tab"""
    pattern_frame = ttk.Frame(notebook)
    notebook.add(pattern_frame, text="Patterns")

    # Create scrollable frame
    canvas = tk.Canvas(pattern_frame, bg=self.colors['bg_primary'])
    scrollbar = ttk.Scrollbar(pattern_frame, orient='vertical', command=canvas.yview)
    scrollable_frame = ttk.Frame(canvas)

    scrollable_frame.bind(
        "<Configure>",
        lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
    )

    canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
    canvas.configure(yscrollcommand=scrollbar.set)

    # Pattern indicators
    detailed_results = result.get('detailed_results', {})
    if isinstance(detailed_results, str):
        try:
            detailed_results = json.loads(detailed_results)
        except Exception:
            detailed_results = {}

    patterns = detailed_results.get('patterns', {})

    for pattern_name, pattern_data in patterns.items():
        self.create_pattern_indicator(scrollable_frame, pattern_name, pattern_data)

    # If no patterns, show default
    if not patterns:
        ttk.Label(scrollable_frame, text="No specific patterns detected",
                 style='Subtitle.TLabel').pack(pady=20)

    canvas.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")


def create_pattern_indicator(self, parent, pattern_name, pattern_data):
    """Create individual pattern indicator"""
    pattern_frame = ttk.Frame(parent, style='Card.TFrame')
    pattern_frame.pack(fill='x', padx=5, pady=5)

    # Pattern header
    header_frame = ttk.Frame(pattern_frame)
    header_frame.pack(fill='x', padx=10, pady=10)

    ttk.Label(header_frame, text=pattern_name.replace('_', ' ').title(),
             font=('Segoe UI', 10, 'bold')).pack(side='left')

    # Score and bar
    if isinstance(pattern_data, dict):
        score = pattern_data.get('score', 0)
        matches = pattern_data.get('matches', 0)
    else:
        score = pattern_data if isinstance(pattern_data, (int, float)) else 0
        matches = 0

    score_label = ttk.Label(header_frame, text=f"{score:.1%}")
    score_label.pack(side='right')

    # Progress bar
    progress_frame = ttk.Frame(pattern_frame)
    progress_frame.pack(fill='x', padx=10, pady=(0, 10))

    progress = ttk.Progressbar(progress_frame, mode='determinate', value=score*100)
    progress.pack(fill='x')

    # Additional info
    if matches > 0:
        ttk.Label(pattern_frame, text=f"Matches found: {matches}",
                 style='Subtitle.TLabel').pack(anchor='w', padx=10, pady=(0, 10))


def create_sentence_tab(self, notebook, result):
    """Create sentence analysis tab"""
    sentence_frame = ttk.Frame(notebook)
    notebook.add(sentence_frame, text="Sentences")

    # Sentence analysis display
    sentence_analysis = result.get('sentence_analysis', {})
    sentences = sentence_analysis.get('sentences', [])

    if sentences:
        # Create text widget with highlighting
        text_widget = scrolledtext.ScrolledText(
            sentence_frame,
            wrap=tk.WORD,
            bg=self.colors['bg_secondary'],
            fg=self.colors['text_primary'],
            selectbackground=self.colors['accent'],
            state='disabled'
        )
        text_widget.pack(fill='both', expand=True, padx=10, pady=10)

        # Configure tags for highlighting
        text_widget.tag_configure('high_risk', background=self.colors['danger'], foreground='white')
        text_widget.tag_configure('medium_risk', background=self.colors['warning'], foreground='white')
        text_widget.tag_configure('low_risk', background=self.colors['success'], foreground='white')

        # Insert sentences with highlighting
        text_widget.config(state='normal')
        for i, sentence_data in enumerate(sentences):
            sentence_text = sentence_data.get('text', '')
            ai_score = sentence_data.get('ai_score', 0)

            start_index = text_widget.index(tk.INSERT)
            text_widget.insert(tk.INSERT, sentence_text + ' ')
            end_index = text_widget.index(tk.INSERT)

            # Apply highlighting based on score
            if ai_score > 0.7:
                text_widget.tag_add('high_risk', start_index, end_index)
            elif ai_score > 0.4:
                text_widget.tag_add('medium_risk', start_index, end_index)
            elif ai_score > 0.2:
                text_widget.tag_add('low_risk', start_index, end_index)

        text_widget.config(state='disabled')

        # Legend
        legend_frame = ttk.Frame(sentence_frame)
        legend_frame.pack(fill='x', padx=10, pady=(0, 10))

        ttk.Label(legend_frame, text="Legend:").pack(side='left')

        for color, label in [('high_risk', 'High Risk'), ('medium_risk', 'Medium Risk'), ('low_risk', 'Low Risk')]:
            legend_item = tk.Label(legend_frame, text=f"● {label}", fg=self.colors['danger'] if color == 'high_risk'
                                 else self.colors['warning'] if color == 'medium_risk' else self.colors['success'],
                                 bg=self.colors['bg_primary'])
            legend_item.pack(side='left', padx=(10, 0))
    else:
        ttk.Label(sentence_frame, text="Sentence analysis not available",
                 style='Subtitle.TLabel').pack(expand=True)


def create_advanced_analysis_tab(self, notebook, result):
    """Create advanced analysis tab"""
    advanced_frame = ttk.Frame(notebook)
    notebook.add(advanced_frame, text="Advanced")

    # Advanced analysis results
    advanced_analyses = result.get('advanced_analyses', {})

    if advanced_analyses:
        # Create scrollable frame
        canvas = tk.Canvas(advanced_frame, bg=self.colors['bg_primary'])
        scrollbar = ttk.Scrollbar(advanced_frame, orient='vertical', command=canvas.yview)
        scrollable_frame = ttk.Frame(canvas)

        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )

        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        # Display each advanced analysis
        for analysis_name, analysis_data in advanced_analyses.items():
            self.create_advanced_analysis_section(scrollable_frame, analysis_name, analysis_data)

        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
    else:
        ttk.Label(advanced_frame, text="Advanced analysis not available",
                 style='Subtitle.TLabel').pack(expand=True)


def create_advanced_analysis_section(self, parent, name, data):
    """Create section for advanced analysis"""
    section_frame = ttk.LabelFrame(parent, text=name.replace('_', ' ').title(), padding=10)
    section_frame.pack(fill='x', padx=5, pady=5)

    if isinstance(data, dict):
        for key, value in data.items():
            if key in ['score', 'confidence'] and isinstance(value, (int, float)):
                ttk.Label(section_frame, text=f"{key.title()}: {value:.1%}").pack(anchor='w')
            elif isinstance(value, (str, int, float, bool)):
                ttk.Label(section_frame, text=f"{key.replace('_', ' ').title()}: {value}").pack(anchor='w')
    else:
        ttk.Label(section_frame, text=str(data)).pack(anchor='w')


def create_recommendations(self, parent, result):
    """Create recommendations section"""
    rec_frame = ttk.Frame(parent, style='Card.TFrame')
    rec_frame.pack(fill='x')

    ttk.Label(rec_frame, text="Recommendations",
             style='Title.TLabel').pack(anchor='w', padx=15, pady=(15, 10))

    # Get recommendations
    recommendations = result.get('recommendations', [])

    if not recommendations:
        # Generate basic recommendations based on score
        ai_score = result.get('ai_score', 0)
        if ai_score > 0.8:
            recommendations = [
                "High AI probability detected. Consider further investigation.",
                "Interview the student about their writing process.",
                "Review assignment submission timeline and behavior."
            ]
        elif ai_score > 0.6:
            recommendations = [
                "Moderate AI indicators present. Consider follow-up.",
                "Provide feedback on academic writing best practices."
            ]
        else:
            recommendations = [
                "Low AI probability. Text appears to be human-written.",
                "Continue monitoring future submissions."
            ]

    # Display recommendations
    for i, rec in enumerate(recommendations, 1):
        rec_item_frame = ttk.Frame(rec_frame)
        rec_item_frame.pack(fill='x', padx=15, pady=2)

        ttk.Label(rec_item_frame, text=f"{i}.", style='Subtitle.TLabel').pack(side='left')
        ttk.Label(rec_item_frame, text=rec, style='Subtitle.TLabel').pack(side='left', padx=(5, 0))

    # Add padding at bottom
    ttk.Label(rec_frame, text="").pack(pady=10)


def analysis_error(self, error_message):
    """Handle analysis error"""
    self.analyze_btn.config(state='normal', text='🔍 Analyze Text')
    self.show_progress(False)
    self.update_status("Analysis failed")
    messagebox.showerror("Analysis Error", f"Failed to analyze text:\n{error_message}")


def get_risk_color(self, score):
    """Get risk indicator color"""
    if score >= 0.8:
        return self.colors['danger']
    elif score >= 0.6:
        return self.colors['warning']
    elif score >= 0.4:
        return '#FFA500'  # Orange
    else:
        return self.colors['success']


def get_risk_text(self, score):
    """Get risk level text"""
    if score >= 0.8:
        return "High"
    elif score >= 0.6:
        return "Medium"
    elif score >= 0.4:
        return "Low"
    else:
        return "Very Low"


def update_word_count(self, event=None):
    """Update word count display"""
    text = self.text_input.get('1.0', tk.END)
    word_count = len(text.split())
    char_count = len(text.strip())
    self.word_count_label.config(text=f"Words: {word_count} | Characters: {char_count}")


