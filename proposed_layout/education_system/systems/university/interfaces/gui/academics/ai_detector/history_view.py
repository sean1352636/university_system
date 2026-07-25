import json
import os
import threading
import time
import random
from datetime import datetime

import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog, scrolledtext

from education_system.systems.university.infrastructure.database.db import DEFAULT_DB_PATH, sqlite3
from education_system.systems.university.infrastructure.auth import UserAuth
from education_system.systems.university.infrastructure.shared_context import get_auth

try:
    from education_system.systems.university.infrastructure.ai.ai_detector.detector import AIDetector
    _AI_DETECTOR_IMPORT_ERROR = None
except Exception as import_error:
    AIDetector = None
    _AI_DETECTOR_IMPORT_ERROR = import_error

try:
    import textract
    TEXTRACT_AVAILABLE = True
except ImportError:
    TEXTRACT_AVAILABLE = False

try:
    from pypdf import PdfReader
    PYPDF2_AVAILABLE = True
except ImportError:
    PYPDF2_AVAILABLE = False

try:
    import docx
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False

from education_system.systems.university.infrastructure.i18n import get_text, _

def create_history_view(self, parent):
    """Create submission history tab"""
    history_frame = ttk.Frame(parent)

    history_frame.pack(fill="both", expand=True)

    # History card
    history_card = ttk.Frame(history_frame, style='Card.TFrame')
    history_card.pack(fill='both', expand=True, padx=10, pady=10)

    # Title and controls
    header_frame = ttk.Frame(history_card)
    header_frame.pack(fill='x', padx=15, pady=15)

    ttk.Label(header_frame, text=get_text("ai_detector.history.title"), style='Title.TLabel').pack(side='left')

    # Refresh button
    self.refresh_history_btn = ttk.Button(header_frame, text=get_text("common.refresh"),
                                        command=self.refresh_history)
    self.refresh_history_btn.pack(side='right')

    # Filter controls
    filter_frame = ttk.Frame(history_card)
    filter_frame.pack(fill='x', padx=15, pady=(0, 15))

    ttk.Label(filter_frame, text=get_text("ai_detector.history.filter_by_student")).pack(side='left')
    self.filter_student_var = tk.StringVar()
    self.filter_entry = ttk.Entry(filter_frame, textvariable=self.filter_student_var, width=20)
    self.filter_entry.pack(side='left', padx=(5, 10))
    self.filter_entry.bind('<Return>', lambda e: self.refresh_history())

    ttk.Button(filter_frame, text=get_text("common.filter"), command=self.refresh_history).pack(side='left')
    ttk.Button(filter_frame, text=get_text("common.clear"), command=self.clear_filter).pack(side='left', padx=(5, 0))

    # History treeview
    tree_frame = ttk.Frame(history_card)
    tree_frame.pack(fill='both', expand=True, padx=15, pady=(0, 15))

    columns = ('ID', 'Student ID', 'Title', 'Course', 'Date', 'AI Score', 'Status')
    column_headers = (
        get_text("common.id"),
        get_text("common.user_id"),
        get_text("table_headers.title"),
        get_text("ai_detector.history.col_course"),
        get_text("common.date"),
        get_text("ai_detector.history.col_ai_score"),
        get_text("common.status")
    )
    self.history_tree = ttk.Treeview(tree_frame, columns=columns, show='headings', height=15)

    # Configure columns
    for col, header in zip(columns, column_headers):
        self.history_tree.heading(col, text=header)
        if col == 'AI Score':
            self.history_tree.column(col, width=100, anchor='center')
        elif col == 'Status':
            self.history_tree.column(col, width=100, anchor='center')
        elif col == 'ID':
            self.history_tree.column(col, width=50, anchor='center')
        else:
            self.history_tree.column(col, width=150)

    # Scrollbars
    v_scrollbar = ttk.Scrollbar(tree_frame, orient='vertical', command=self.history_tree.yview)
    h_scrollbar = ttk.Scrollbar(tree_frame, orient='horizontal', command=self.history_tree.xview)
    self.history_tree.configure(yscrollcommand=v_scrollbar.set, xscrollcommand=h_scrollbar.set)

    # Pack treeview and scrollbars
    self.history_tree.grid(row=0, column=0, sticky='nsew')
    v_scrollbar.grid(row=0, column=1, sticky='ns')
    h_scrollbar.grid(row=1, column=0, sticky='ew')

    tree_frame.grid_rowconfigure(0, weight=1)
    tree_frame.grid_columnconfigure(0, weight=1)

    # Bind double-click event
    self.history_tree.bind('<Double-1>', self.view_submission_details)

    # Load initial history
    self.refresh_history()


def refresh_history(self):
    """Refresh submission history"""
    if not hasattr(self, 'history_tree'):
        return  # History view not created yet
    try:
        # Get filter (filter_student_var only exists after history view is created)
        student_filter = None
        if hasattr(self, 'filter_student_var') and self.filter_student_var.get().strip():
            student_filter = self.filter_student_var.get()

        # Get submissions
        history_data = self.detector.get_submission_history(student_id=student_filter, limit=100)
        submissions = history_data.get('submissions', [])

        # Clear existing items
        for item in self.history_tree.get_children():
            self.history_tree.delete(item)

        # Populate tree
        for submission in submissions:
            submission_id = submission.get('id', '')
            student_id = submission.get('student_id', '')
            title = submission.get('title', 'Untitled')
            course = submission.get('course_code', '')
            date = submission.get('submission_date', '')
            ai_score = submission.get('ai_score', 0)
            is_ai = submission.get('is_ai_generated', False)

            # Format date
            try:
                if date:
                    date_obj = datetime.fromisoformat(date.replace('Z', '+00:00'))
                    formatted_date = date_obj.strftime('%Y-%m-%d %H:%M')
                else:
                    formatted_date = 'Unknown'
            except Exception:
                formatted_date = date

            # Format score
            score_text = f"{ai_score:.1%}" if ai_score is not None else "N/A"

            # Status
            status = get_text("ai_detector.history.status_ai") if is_ai else get_text("ai_detector.history.status_human")

            self.history_tree.insert('', 'end', values=(
                submission_id, student_id, title, course, formatted_date, score_text, status
            ))

        self.update_status(f"Loaded {len(submissions)} submissions")

    except Exception as e:
        messagebox.showerror(get_text("common.error"), get_text("ai_detector.history.load_failed", error=str(e)))


def clear_filter(self):
    """Clear history filter"""
    self.filter_student_var.set("")
    self.refresh_history()


def view_submission_details(self, event):
    """View detailed submission information"""
    selection = self.history_tree.selection()
    if not selection:
        return

    item = self.history_tree.item(selection[0])
    submission_id = item['values'][0]

    try:
        submission_details = self.detector.get_submission_details(submission_id)
        self.show_submission_details_window(submission_details)
    except Exception as e:
        messagebox.showerror(get_text("common.error"), get_text("ai_detector.history.details_failed", error=str(e)))


def show_submission_details_window(self, details):
    """Show submission details in new window"""
    details_window = tk.Toplevel(self.root)
    details_window.title(get_text("ai_detector.history.details_title"))
    details_window.geometry("800x600")
    details_window.configure(bg=self.colors['bg_primary'])

    # Create scrollable frame
    canvas = tk.Canvas(details_window, bg=self.colors['bg_primary'])
    scrollbar = ttk.Scrollbar(details_window, orient='vertical', command=canvas.yview)
    scrollable_frame = ttk.Frame(canvas)

    scrollable_frame.bind(
        "<Configure>",
        lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
    )

    canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
    canvas.configure(yscrollcommand=scrollbar.set)

    # Content
    ttk.Label(scrollable_frame, text=get_text("ai_detector.history.details_title"),
             style='Title.TLabel').pack(anchor='w', padx=20, pady=20)

    # Basic info
    info_frame = ttk.LabelFrame(scrollable_frame, text=get_text("ai_detector.history.basic_info"), padding=15)
    info_frame.pack(fill='x', padx=20, pady=(0, 20))

    for key, value in details.items():
        if key not in ['submission_text', 'detailed_results'] and value is not None:
            ttk.Label(info_frame, text=f"{key.replace('_', ' ').title()}: {value}").pack(anchor='w')

    # Text content
    if details.get('submission_text'):
        text_frame = ttk.LabelFrame(scrollable_frame, text=get_text("ai_detector.history.submission_text"), padding=15)
        text_frame.pack(fill='both', expand=True, padx=20, pady=(0, 20))

        text_widget = scrolledtext.ScrolledText(
            text_frame, wrap=tk.WORD, height=15,
            bg=self.colors['bg_secondary'], fg=self.colors['text_primary']
        )
        text_widget.pack(fill='both', expand=True)
        text_widget.insert('1.0', details['submission_text'])
        text_widget.config(state='disabled')

    canvas.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")


